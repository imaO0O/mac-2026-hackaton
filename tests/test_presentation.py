"""Презентация собирается из отчётов, а не набирается руками.

Генератор (`scripts/make_presentation.py`) берёт числа теми же функциями, что сверяют
README, а вердикты «включено / отклонено» — из отчётов проверок. Тесты стерегут
связку: переименовали отчёт проверки — слайд решений молча потерял бы строку;
переписали пример карточки в README — слайд карточки вышел бы пустым.
"""
from __future__ import annotations

import pytest

pytest.importorskip("pptx")

from nefte.config import ROOT  # noqa: E402

# Проверки 21.09 ещё без отчёта в части сборок: их строка появляется с отчётом.
PENDING = {"severity_veto.json", "t95_conservative.json", "aging_bounds.json"}


def test_every_decision_points_to_a_check_that_writes_it():
    from scripts.make_presentation import DECISIONS

    scripts = "\n".join(p.read_text(encoding="utf-8")
                        for p in (ROOT / "scripts").glob("*.py"))
    for label, name, _ in DECISIONS:
        assert label.strip(), name
        assert name in scripts, f"{name}: ни один скрипт его не пишет"
        if name not in PENDING and not (ROOT / "reports" / name).exists():
            pytest.skip(f"нет reports/{name} — отчёты не собраны")


def test_card_slide_takes_the_readme_example():
    from scripts.make_presentation import card_example

    cmd, block = card_example()
    assert cmd.startswith("python scripts/run_cycle.py")
    assert "Действие:" in block and "Что дальше:" in block


def test_deck_builds_from_reports(tmp_path, monkeypatch):
    import scripts.make_presentation as mp

    if not (ROOT / "reports" / "architectures.json").exists():
        pytest.skip("отчёты не собраны")
    monkeypatch.setattr(mp, "OUT", tmp_path / "deck.pptx")
    assert mp.main() == 0
    from pptx import Presentation

    deck = Presentation(str(tmp_path / "deck.pptx"))
    assert len(deck.slides) == 14
    for slide in deck.slides:
        assert slide.has_notes_slide or slide == deck.slides[0]
